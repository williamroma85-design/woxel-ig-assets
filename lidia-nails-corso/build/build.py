import os, sys
sys.path.insert(0, os.path.dirname(__file__))
from contenuti import *
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, ListFlowable, ListItem, KeepTogether
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT = os.path.join(os.path.dirname(__file__), '..', 'output'); os.makedirs(OUT, exist_ok=True)
ROSA = colors.HexColor('#B76E79'); NERO = colors.HexColor('#2B2B2B'); CIPRIA = colors.HexColor('#F6ECE9')

def pdf(path, titolo, sotto, blocchi):
    h1 = ParagraphStyle('h1', fontName='Helvetica-Bold', fontSize=26, textColor=ROSA, leading=30, spaceAfter=6)
    h2 = ParagraphStyle('h2', fontName='Helvetica-Bold', fontSize=14, textColor=NERO, spaceBefore=14, spaceAfter=6)
    sub = ParagraphStyle('sub', fontName='Helvetica', fontSize=12, textColor=colors.grey, spaceAfter=18)
    body = ParagraphStyle('b', fontName='Helvetica', fontSize=10.5, leading=15, textColor=NERO)
    def foot(c, d):
        c.setFont('Helvetica', 8); c.setFillColor(colors.grey)
        c.drawString(2*cm, 1.2*cm, BRAND); c.drawRightString(19*cm, 1.2*cm, str(d.page))
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm, title=titolo, author=BRAND)
    st = [Paragraph(BRAND.upper(), ParagraphStyle('k', fontName='Helvetica-Bold', fontSize=10, textColor=ROSA, spaceAfter=10)),
          Paragraph(titolo, h1), Paragraph(sotto, sub)]
    for t, pts in blocchi:
        items = [ListItem(Paragraph(p.replace('&','&amp;'), body), leftIndent=14, bulletColor=ROSA) for p in pts]
        st.append(KeepTogether([Paragraph(t, h2), ListFlowable(items, bulletType='bullet', bulletFontSize=7, start='•')]))
    doc.build(st, onFirstPage=foot, onLaterPages=foot)

pdf(os.path.join(OUT, 'guida-professionale.pdf'), TITOLO, SOTTOTITOLO, [(t, p) for t, p, _, _ in SEZIONI])
pdf(os.path.join(OUT, 'mini-guida-cliente.pdf'), "La cura delle tue Gel Tip", "Piccola guida per la cliente", MINI_GUIDA)

# Slide: titolo + max 4 punti brevi, note del relatore
def short(s, n=95):
    s = s.split(': ')[0] if len(s) > n and ': ' in s else s
    return s if len(s) <= n else s[:n-1].rsplit(' ', 1)[0] + '…'

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]
def rgb(h): return RGBColor.from_string(h)
def bg(s, h):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height); r.fill.solid(); r.fill.fore_color.rgb = rgb(h); r.line.fill.background()
def txt(s, t, x, y, w, h, size, bold=False, color='2B2B2B'):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(color); r.font.name = 'Calibri'; return tf

s = prs.slides.add_slide(blank); bg(s, 'F6ECE9')
txt(s, BRAND.upper(), 0.9, 2.2, 11, 0.6, 20, True, 'B76E79'); txt(s, TITOLO, 0.9, 2.9, 11.5, 1.5, 48, True); txt(s, SOTTOTITOLO, 0.9, 4.6, 11, 0.6, 22, color='777777')
IMMAGINI = {3: 'struttura-tip', 4: 'tip-tipi', 5: 'misure-mano', 6: 'unghia-naturale', 7: 'scelta-misura', 9: 'applicazione', 10: 'inclinazione', 12: 'rimozione', 14: 'semaforo'}
for idx, ((t, pts, nota, on), brevi) in enumerate(zip(SEZIONI, SLIDE_PUNTI)):
    if not on: continue
    s = prs.slides.add_slide(blank); bg(s, 'FFFFFF')
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.25), prs.slide_height); bar.fill.solid(); bar.fill.fore_color.rgb = rgb('B76E79'); bar.line.fill.background()
    txt(s, t, 0.9, 0.6, 11.5, 1.0, 36, True)
    img = IMMAGINI.get(idx); imgp = os.path.join(OUT, 'immagini', (img or '') + '.png')
    tf = txt(s, '', 0.9, 1.9, 6.1 if img else 11.5, 4.8, 24)
    first = True
    for p in brevi:
        para = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        r = para.add_run(); r.text = '•  ' + p; r.font.size = Pt(24 if img else 28); r.font.name = 'Calibri'; r.font.color.rgb = rgb('2B2B2B'); para.space_after = Pt(18)
    if img and os.path.exists(imgp):
        from PIL import Image as _I; w, h = _I.open(imgp).size; wi = 5.5; hi = wi * h / w
        if hi > 5.0: wi, hi = 5.0 * w / h, 5.0
        s.shapes.add_picture(imgp, Inches(12.9 - wi), Inches(1.7), width=Inches(wi))
    txt(s, BRAND, 0.9, 6.9, 4, 0.4, 12, color='B76E79')
    s.notes_slide.notes_text_frame.text = nota + "\n\nDettaglio:\n" + "\n".join("- " + p for p in pts)
prs.save(os.path.join(OUT, 'slide-lezione.pptx'))
print('ok')
